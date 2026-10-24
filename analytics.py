"""
analytics.py
------------
Core calculation + Excel-generation logic, kept separate from the Flask
routes (app.py) so it can be re-used by the "Update Analysis" button and by
the automatic post-save sync described in the spec (section 8.5 / 9.2).

Unit conversion (from Data_Dictionary, section 8.1 of the spec):
    temp_celsius      = temp * 41
    atemp_celsius     = atemp * 50
    hum_percent       = hum * 100
    windspeed_original= windspeed * 67
These are the formulas used any time a raw/normalized value must be shown
to, or collected from, a website user.
"""

import csv
import sqlite3
import statistics
import datetime
import os
import openpyxl
from openpyxl.utils import get_column_letter

DB_PATH = os.path.join(os.path.dirname(__file__), "bikeshare.db")
TEMPLATE_XLSX = os.path.join(os.path.dirname(__file__), "data", "Bike_Sharing_Analysis_Final.xlsx")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "exports")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ---------- Unit conversion helpers (section 8.1) ----------

def temp_to_celsius(norm_temp):
    return round(norm_temp * 41, 2)

def celsius_to_norm_temp(celsius):
    return round(celsius / 41, 4)

def atemp_to_celsius(norm_atemp):
    return round(norm_atemp * 50, 2)

def celsius_to_norm_atemp(celsius):
    return round(celsius / 50, 4)

def hum_to_percent(norm_hum):
    return round(norm_hum * 100, 1)

def percent_to_norm_hum(pct):
    return round(pct / 100, 4)

def windspeed_to_original(norm_ws):
    return round(norm_ws * 67, 2)

def original_to_norm_windspeed(ws):
    return round(ws / 67, 4)


def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def combined_rows(conn):
    """Raw_Data (historical, protected) + Web_Entries (user-added), combined per spec 8.3."""
    raw = conn.execute("SELECT * FROM raw_data").fetchall()
    web = conn.execute("SELECT * FROM web_entries").fetchall()
    return list(raw) + list(web)


# ---------- Descriptive statistics ----------

def compute_descriptive_stats(conn):
    rows = combined_rows(conn)
    cnts = [r["cnt"] for r in rows]
    if not cnts:
        return {}
    try:
        mode = statistics.mode(cnts)
    except statistics.StatisticsError:
        mode = None
    n = len(cnts)
    mean = statistics.mean(cnts)
    stdev = statistics.stdev(cnts) if n > 1 else 0
    return {
        "count": n,
        "mean": round(mean, 4),
        "median": statistics.median(cnts),
        "mode": mode,
        "std_dev": round(stdev, 4),
        "variance": round(stdev ** 2, 4),
        "min": min(cnts),
        "max": max(cnts),
        "total_rentals": sum(cnts),
    }


# ---------- Correlation (Pearson) ----------

def _pearson(x, y):
    n = len(x)
    if n < 2:
        return 0.0
    mx, my = sum(x) / n, sum(y) / n
    num = sum((xi - mx) * (yi - my) for xi, yi in zip(x, y))
    denx = sum((xi - mx) ** 2 for xi in x) ** 0.5
    deny = sum((yi - my) ** 2 for yi in y) ** 0.5
    if denx == 0 or deny == 0:
        return 0.0
    return num / (denx * deny)


def compute_correlation(conn):
    rows = combined_rows(conn)
    temp = [r["temp"] for r in rows]
    hum = [r["hum"] for r in rows]
    wind = [r["windspeed"] for r in rows]
    cnt = [r["cnt"] for r in rows]
    return {
        "temp_cnt": round(_pearson(temp, cnt), 4),
        "hum_cnt": round(_pearson(hum, cnt), 4),
        "windspeed_cnt": round(_pearson(wind, cnt), 4),
        "temp_hum": round(_pearson(temp, hum), 4),
    }


# ---------- Simple linear regression: cnt ~ temp ----------

def compute_regression(conn):
    rows = combined_rows(conn)
    x = [r["temp"] for r in rows]
    y = [r["cnt"] for r in rows]
    n = len(x)
    if n < 2:
        return {}
    mx, my = sum(x) / n, sum(y) / n
    sxx = sum((xi - mx) ** 2 for xi in x)
    sxy = sum((xi - mx) * (yi - my) for xi, yi in zip(x, y))
    slope = sxy / sxx if sxx else 0
    intercept = my - slope * mx
    r = _pearson(x, y)
    r_sq = r ** 2
    return {
        "slope": round(slope, 4),
        "intercept": round(intercept, 4),
        "r_value": round(r, 4),
        "r_squared": round(r_sq, 4),
        "equation": f"cnt = {round(intercept, 2)} + {round(slope, 2)} * temp",
        "observations": n,
    }


# ---------- Pivot-style summaries ----------

def compute_pivot(conn, field):
    rows = combined_rows(conn)
    totals = {}
    for r in rows:
        key = r[field]
        totals[key] = totals.get(key, 0) + r["cnt"]
    return dict(sorted(totals.items()))


def compute_monthly_summary(conn):
    rows = combined_rows(conn)
    monthly = {}
    for r in rows:
        key = (r["yr"], r["mnth"])
        if key not in monthly:
            monthly[key] = {"total": 0, "casual": 0, "registered": 0, "records": 0}
        monthly[key]["total"] += r["cnt"]
        monthly[key]["casual"] += r["casual"]
        monthly[key]["registered"] += r["registered"]
        monthly[key]["records"] += 1
    out = []
    for (yr, mnth), v in sorted(monthly.items()):
        year = 2011 + yr
        avg = round(v["total"] / v["records"], 3) if v["records"] else 0
        out.append({"year": year, "month": mnth, **v, "avg_per_record": avg})
    return out


def run_full_analysis():
    """Recomputes every analysis result. Called after each add/edit/delete and
    by the manual 'Update Analysis' button (spec 6.3 / 8.5)."""
    conn = get_conn()
    try:
        result = {
            "descriptive": compute_descriptive_stats(conn),
            "correlation": compute_correlation(conn),
            "regression": compute_regression(conn),
            "pivot_hourly": compute_pivot(conn, "hr"),
            "pivot_season": compute_pivot(conn, "season"),
            "pivot_weather": compute_pivot(conn, "weathersit"),
            "monthly": compute_monthly_summary(conn),
            "computed_at": datetime.datetime.now().isoformat(timespec="seconds"),
        }
        return result
    finally:
        conn.close()


# ---------- Excel export (section 7 / 9) ----------

def export_latest_data():
    """7.1: export current DB records (raw_data + web_entries) to
    Bike_Sharing_Export.xlsx. Does NOT touch the original workbook."""
    conn = get_conn()
    rows = combined_rows(conn)
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Exported_Data"
    headers = ["instant", "dteday", "season", "yr", "mnth", "hr", "holiday", "weekday",
               "workingday", "weathersit", "temp", "atemp", "hum", "windspeed",
               "casual", "registered", "cnt"]
    ws.append(headers)
    for r in rows:
        ws.append([r[h] for h in headers])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(h) + 2)

    out_path = os.path.join(OUTPUT_DIR, "Bike_Sharing_Export.xlsx")
    wb.save(out_path)
    return out_path, len(rows)


def generate_updated_analysis_workbook():
    """7.2 / 9.2: use the ORIGINAL workbook as a template, insert latest data
    and recalculated results into a NEW file. Original is never modified."""
    wb = openpyxl.load_workbook(TEMPLATE_XLSX)

    # Step 1: update/replace Web_Entries with current DB web_entries rows
    conn = get_conn()
    web_rows = conn.execute("SELECT * FROM web_entries").fetchall()
    headers = ["instant", "dteday", "season", "yr", "mnth", "hr", "holiday", "weekday",
               "workingday", "weathersit", "temp", "atemp", "hum", "windspeed",
               "casual", "registered", "cnt"]
    ws_web = wb["Web_Entries"]
    # clear old data rows (keep header row 1)
    if ws_web.max_row > 1:
        ws_web.delete_rows(2, ws_web.max_row - 1)
    for r in web_rows:
        ws_web.append([r[h] for h in headers])

    # Step 2 & 3: recalculate descriptive / correlation / regression
    analysis = run_full_analysis()

    ws_desc = wb["Descriptive_Result"]
    desc = analysis["descriptive"]
    desc_map = {
        "Mean": desc.get("mean"), "Median": desc.get("median"), "Mode": desc.get("mode"),
        "Standard Deviation": desc.get("std_dev"), "Sample Variance": desc.get("variance"),
    }
    for row in range(1, ws_desc.max_row + 1):
        label = ws_desc.cell(row, 1).value
        if label in desc_map and desc_map[label] is not None:
            ws_desc.cell(row, 2).value = desc_map[label]

    corr = analysis["correlation"]
    ws_corr = wb["Correlation_Result"]
    # cnt row intersect temp/hum/windspeed columns (B2:D2 header temp,hum,windspeed; row5=cnt)
    ws_corr.cell(5, 2).value = corr["temp_cnt"]
    ws_corr.cell(5, 3).value = corr["hum_cnt"]
    ws_corr.cell(5, 4).value = corr["windspeed_cnt"]

    # Step 4: refresh Tableau_Data monthly summary + pivot-style sheets
    ws_tab = wb["Tableau_Data"]
    if ws_tab.max_row > 4:
        ws_tab.delete_rows(5, ws_tab.max_row - 4)
    for m in analysis["monthly"]:
        ws_tab.append([
            m["year"], m["month"],
            datetime.date(m["year"], m["month"], 1),
            m["total"], m["casual"], m["registered"], m["records"], m["avg_per_record"],
        ])

    def refresh_pivot(sheet_name, data_dict):
        ws = wb[sheet_name]
        if ws.max_row > 3:
            ws.delete_rows(4, ws.max_row - 3)
        for k, v in data_dict.items():
            ws.append([k, v])
        ws.append(["Grand Total", sum(data_dict.values())])

    refresh_pivot("Pivot_Hourly", analysis["pivot_hourly"])
    refresh_pivot("Pivot_Season", analysis["pivot_season"])
    refresh_pivot("Pivot_Weather", analysis["pivot_weather"])

    conn.close()

    out_path = os.path.join(OUTPUT_DIR, "Bike_Sharing_Analysis_Updated.xlsx")
    wb.save(out_path)
    return out_path, analysis


# ---------- Tableau Public data files (CSV + Excel) ----------

TABLEAU_COLUMNS = ["instant", "dteday", "season", "yr", "mnth", "hr", "holiday", "weekday",
                   "workingday", "weathersit", "temp", "atemp", "hum", "windspeed",
                   "casual", "registered", "cnt"]


def _tableau_rows(conn):
    """All records currently in the database, in the original dataset's 17-column
    structure. Original rows keep their instant; website rows are numbered on from the
    last original instant (their own ids restart at 1 and would otherwise clash), and
    every date is written as 'YYYY-MM-DD 00:00:00' like the original dataset."""
    raw = conn.execute("SELECT * FROM raw_data ORDER BY instant").fetchall()
    web = conn.execute("SELECT * FROM web_entries ORDER BY instant").fetchall()
    out = []
    for r in raw:
        row = [r[c] for c in TABLEAU_COLUMNS]
        row[1] = str(row[1])[:10] + " 00:00:00"
        out.append(row)
    next_id = (max(r["instant"] for r in raw) if raw else 0) + 1
    for r in web:
        row = [r[c] for c in TABLEAU_COLUMNS]
        row[0] = next_id
        next_id += 1
        row[1] = str(row[1])[:10] + " 00:00:00"
        out.append(row)
    return out


def export_tableau_csv():
    """CSV of the current database for updating the Tableau Public data source."""
    conn = get_conn()
    try:
        rows = _tableau_rows(conn)
    finally:
        conn.close()
    out_path = os.path.join(OUTPUT_DIR, "Bike_Sharing_Tableau_Data.csv")
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(TABLEAU_COLUMNS)
        w.writerows(rows)
    return out_path, len(rows)


def export_tableau_xlsx():
    """Excel version of the same Tableau data file."""
    conn = get_conn()
    try:
        rows = _tableau_rows(conn)
    finally:
        conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tableau_Data"
    ws.append(TABLEAU_COLUMNS)
    for row in rows:
        row = list(row)
        row[1] = datetime.datetime.strptime(row[1][:10], "%Y-%m-%d")
        ws.append(row)
    for i, h in enumerate(TABLEAU_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 20 if h == "dteday" else max(10, len(h) + 2)
    for cell in ws["B"][1:]:
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    out_path = os.path.join(OUTPUT_DIR, "Bike_Sharing_Tableau_Data.xlsx")
    tmp_path = out_path + ".tmp"
    wb.save(tmp_path)
    os.replace(tmp_path, out_path)      # atomic: readers never see a half-written file
    return out_path, len(rows)
