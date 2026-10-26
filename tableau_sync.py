"""
tableau_sync.py
---------------
Dashboard update pipeline:

    (a) Excel file  ->  .hyper extract      (Tableau Hyper API)
    (b) .hyper      ->  .twb workbook       (the workbook's data source is re-pointed to the .hyper)
    (c) .twb + .hyper are zipped into a .twbx for publishing to Tableau Cloud.

Files (all under static/tableau/):
    template/Project_data_Analysis.twb   original workbook from Tableau (never modified)
    Bike_Sharing.hyper                   generated extract
    Project_data_Analysis.twb            generated workbook, connected to the .hyper
    Project_data_Analysis.twbx           generated packaged workbook (twb + hyper) for publishing

Run it by hand too:
    python tableau_sync.py                       # uses exports/Bike_Sharing_Tableau_Data.xlsx
    python tableau_sync.py --excel my_file.xlsx  # any Excel file with the 17 dataset columns
"""

import argparse
import copy
import datetime
import os
import shutil
import zipfile

from dotenv import load_dotenv
from lxml import etree
from tableauhyperapi import (
    Connection, CreateMode, HyperProcess, Inserter, SqlType, TableDefinition,
    TableName, Telemetry,
)

load_dotenv()

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TABLEAU_DIR = os.path.join(BASE_DIR, "static", "tableau")
TEMPLATE_TWB = os.path.join(TABLEAU_DIR, "template", "Project_data_Analysis.twb")
HYPER_PATH = os.path.join(TABLEAU_DIR, "Bike_Sharing.hyper")
APP_HYPER_PATH = os.path.join(BASE_DIR, "data", "Bike_Sharing.hyper")
TWB_PATH = os.path.join(TABLEAU_DIR, "Project_data_Analysis.twb")
TWBX_PATH = os.path.join(TABLEAU_DIR, "Project_data_Analysis.twbx")
DEFAULT_EXCEL = os.path.join(BASE_DIR, "exports", "Bike_Sharing_Tableau_Data.xlsx")

HYPER_TABLE = TableName("Extract", "Extract")   # same schema/table name Tableau itself uses
HYPER_ARCHIVE_PATH = "Data/Extracts/Bike_Sharing.hyper"   # location of the .hyper inside the .twbx

# Column name -> Hyper type. The workbook uses all of these (including yr, mnth and atemp).
COLUMNS = [
    ("instant", SqlType.big_int()), ("dteday", SqlType.date()), ("season", SqlType.big_int()),
    ("yr", SqlType.big_int()), ("mnth", SqlType.big_int()), ("hr", SqlType.big_int()),
    ("holiday", SqlType.big_int()), ("weekday", SqlType.big_int()), ("workingday", SqlType.big_int()),
    ("weathersit", SqlType.big_int()), ("temp", SqlType.double()), ("atemp", SqlType.double()),
    ("hum", SqlType.double()), ("windspeed", SqlType.double()), ("casual", SqlType.big_int()),
    ("registered", SqlType.big_int()), ("cnt", SqlType.big_int()),
]
INT_COLS = {n for n, t in COLUMNS if t == SqlType.big_int()}
FLOAT_COLS = {"temp", "atemp", "hum", "windspeed"}


# ---------------------------------------------------------------- (a) Excel -> .hyper

def _read_excel_rows(excel_path):
    """Read the dataset rows from the Excel file. Uses the 'Tableau_Data' sheet when it exists,
    otherwise the first sheet. Blank rows are skipped (they would otherwise become all-NULL rows)."""
    import openpyxl

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb["Tableau_Data"] if "Tableau_Data" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    missing = [n for n, _ in COLUMNS if n not in header]
    if missing:
        raise ValueError(f"Excel file is missing column(s): {', '.join(missing)}")
    idx = [header.index(n) for n, _ in COLUMNS]

    rows = []
    for raw in it:
        vals = [raw[i] if i < len(raw) else None for i in idx]
        if vals[0] is None:          # no 'instant' -> empty row
            continue
        row = []
        for (name, _), v in zip(COLUMNS, vals):
            if name == "dteday":
                if isinstance(v, datetime.datetime):
                    v = v.date()
                elif isinstance(v, str):
                    v = datetime.date.fromisoformat(v[:10])
            elif name in INT_COLS:
                v = int(v)
            elif name in FLOAT_COLS:
                v = float(v)
            row.append(v)
        rows.append(row)
    wb.close()
    return rows


def excel_to_hyper(excel_path=DEFAULT_EXCEL, hyper_path=HYPER_PATH):
    """Step (a): rebuild the .hyper extract from the Excel file. Returns the row count."""
    from tableauhyperapi import Date
    rows = _read_excel_rows(excel_path)
    for r in rows:
        d = r[1]
        r[1] = Date(d.year, d.month, d.day)

    os.makedirs(os.path.dirname(hyper_path), exist_ok=True)
    table = TableDefinition(HYPER_TABLE, [TableDefinition.Column(n, t) for n, t in COLUMNS])
    with HyperProcess(Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU, parameters={"log_config": ""}) as hyper:
        with Connection(hyper.endpoint, hyper_path, CreateMode.CREATE_AND_REPLACE) as conn:
            conn.catalog.create_schema_if_not_exists(HYPER_TABLE.schema_name)
            conn.catalog.create_table(table)
            with Inserter(conn, table) as ins:
                ins.add_rows(rows)
                ins.execute()
            return conn.execute_scalar_query(f"SELECT COUNT(*) FROM {HYPER_TABLE}")


# ---------------------------------------------------------------- (b) .hyper -> .twb

def _federated_datasource(root):
    for ds in root.iter("datasource"):
        conn = ds.find("connection")
        if ds.get("name", "").startswith("federated") and conn is not None and conn.get("class") == "federated":
            return ds
    raise ValueError("No federated data source found in the workbook.")


def patch_twb(template_twb=TEMPLATE_TWB, out_twb=TWB_PATH, hyper_ref=HYPER_ARCHIVE_PATH):
    """Step (b): copy the template workbook and point its data source at the .hyper file.

    The original data source is a union of the Excel sheets Raw_Data + Web_Entries. It is replaced by a
    single table of the .hyper, which already contains both. Every field the charts use keeps its name,
    so worksheets, calculated fields, parameters and the dashboard layout are left untouched."""
    if not os.path.isfile(template_twb):
        # Template missing (e.g. the folder was copied without static/tableau/template/).
        # An already generated workbook is connected to the .hyper, which is exactly what we need,
        # because the .hyper is rebuilt in place and the .twb only references it by path.
        # BUT if that workbook is still the Excel-connected original, use it as the template.
        if os.path.isfile(out_twb):
            if "excel-direct" not in open(out_twb, encoding="utf-8").read():
                return out_twb
            os.makedirs(os.path.dirname(template_twb), exist_ok=True)
            shutil.copyfile(out_twb, template_twb)
        else:
            raise FileNotFoundError(
            "Template workbook not found. Put your original Tableau workbook at "
            f"{template_twb} (it must be the .twb saved from Tableau, connected to the Excel file).")
    tree = etree.parse(template_twb)
    root = tree.getroot()
    ds = _federated_datasource(root)
    conn = ds.find("connection")

    # 1. connection: Excel (union of 2 sheets) -> .hyper (1 table)
    conn_id = "hyper.bikesharing0extract0000000000001"
    old_named = conn.find("named-connections")
    new_named = etree.Element("named-connections")
    nc = etree.SubElement(new_named, "named-connection", caption="Bike_Sharing", name=conn_id)
    etree.SubElement(nc, "connection", {
        "authentication": "auth-none", "class": "hyper", "dbname": hyper_ref,
        "default-settings": "yes", "schema": "Extract", "sslmode": "", "tablename": "Extract", "username": "",
    })
    conn.replace(old_named, new_named)

    def table_relation():
        return etree.Element("relation", {"connection": conn_id, "name": "Extract",
                                          "table": "[Extract].[Extract]", "type": "table"})
    conn.replace(conn.find("relation"), table_relation())

    # Tableau stores a second extract connection under <extract>. Repoint it too;
    # otherwise Cloud may try to open the stale Hyper path from the source workbook.
    extract = ds.find("extract")
    extract_conn = extract.find("connection") if extract is not None else None
    if extract_conn is not None:
        extract_conn.set("dbname", hyper_ref)
        extract_conn.set("schema", "Extract")
        extract_conn.set("tablename", "Extract")
        extract_relation = extract_conn.find("relation")
        if extract_relation is not None:
            extract_relation.set("name", "Extract")
            extract_relation.set("table", "[Extract].[Extract]")
            extract_relation.set("type", "table")

    # 2. metadata: columns now come from [Extract]; drop the Excel-only helper columns
    keep = {n for n, _ in COLUMNS}
    meta = conn.find("metadata-records")
    for rec in list(meta):
        if rec.get("class") == "column" and rec.findtext("remote-name") not in keep:
            meta.remove(rec)
            continue
        pn = rec.find("parent-name")
        if pn is not None and pn.text == "[Union]":
            pn.text = "[Extract]"
    for i, rec in enumerate(r for r in meta if r.get("class") == "column"):
        rec.find("ordinal").text = str(i)

    # 3. declared columns for the dropped helper fields
    for col in ds.findall("column"):
        n = (col.get("name") or "").strip("[]")
        if n in {"F18", "Sheet", "Table Name"} or n.startswith("คำแนะนำสำหรับผู้พัฒนาเว็บไซต์"):
            ds.remove(col)

    # 4. object graph (the "logical table" Tableau shows on the Data Source tab)
    og = ds.find("object-graph")
    if og is not None:
        for obj in og.iter("object"):
            props = obj.find("properties")
            if props is not None and props.find("relation") is not None:
                props.replace(props.find("relation"), table_relation())
                obj.set("caption", "Bike_Sharing")

    os.makedirs(os.path.dirname(out_twb), exist_ok=True)
    tree.write(out_twb, encoding="utf-8", xml_declaration=True)
    return out_twb


# ---------------------------------------------------------------- (c) package for publishing

def package_twbx(twb_path=TWB_PATH, hyper_path=HYPER_PATH, out_twbx=TWBX_PATH):
    with zipfile.ZipFile(out_twbx, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(twb_path, os.path.basename(twb_path))
        z.write(hyper_path, HYPER_ARCHIVE_PATH)
    return out_twbx


def package_app_hyper(
    hyper_path=APP_HYPER_PATH,
    template_twb=TEMPLATE_TWB,
    twb_path=TWB_PATH,
    out_twbx=TWBX_PATH,
):
    """Package the app's latest Hyper extract with the Tableau workbook template."""
    if not os.path.isfile(hyper_path):
        raise FileNotFoundError(f"App Hyper extract not found: {hyper_path}")
    patch_twb(template_twb, twb_path)
    problems = validate_workbook(twb_path, hyper_path)
    if problems:
        raise RuntimeError("; ".join(problems))
    return package_twbx(twb_path, hyper_path, out_twbx)


def validate_workbook(twb_path=TWB_PATH, hyper_path=HYPER_PATH):
    """Cheap safety checks: every field the worksheets use exists in the .hyper and no Excel link is left."""
    text = open(twb_path, encoding="utf-8").read()
    problems = []
    if "excel-direct" in text:
        problems.append("workbook still references the Excel connection")
    root = etree.fromstring(text.encode("utf-8"))
    ds = _federated_datasource(root)
    hyper_connections = [
        conn for conn in ds.iter("connection") if conn.get("class") == "hyper"
    ]
    if not hyper_connections:
        problems.append("workbook has no Hyper connection")
    for conn in hyper_connections:
        if conn.get("dbname") != HYPER_ARCHIVE_PATH:
            problems.append(
                f"workbook Hyper connection points to {conn.get('dbname')!r}, "
                f"expected {HYPER_ARCHIVE_PATH!r}"
            )
    with HyperProcess(Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU, parameters={"log_config": ""}) as hyper:
        with Connection(hyper.endpoint, hyper_path) as conn:
            cols = {c.name.unescaped for c in conn.catalog.get_table_definition(HYPER_TABLE).columns}
    meta = {r.findtext("remote-name") for r in ds.find("connection/metadata-records") if r.get("class") == "column"}
    if meta - cols:
        problems.append(f"workbook fields not in .hyper: {sorted(meta - cols)}")
    if cols - meta:
        problems.append(f".hyper columns unknown to workbook: {sorted(cols - meta)}")
    return problems


def tableau_cloud_config():
    """Return Tableau Cloud settings when PAT credentials are configured."""
    token_name = os.environ.get("TABLEAU_PAT_NAME")
    token_secret = os.environ.get("TABLEAU_PAT_SECRET")
    if not token_name and not token_secret:
        return None
    if not token_name or not token_secret:
        raise RuntimeError("Set both TABLEAU_PAT_NAME and TABLEAU_PAT_SECRET.")
    return {
        "server_url": os.environ.get(
            "TABLEAU_SERVER_URL", "https://prod-apsoutheast-b.online.tableau.com"
        ).rstrip("/"),
        "site_content_url": os.environ.get(
            "TABLEAU_SITE_CONTENT_URL", "thanatchanunboonkasem-8b976e5748"
        ),
        "token_name": token_name,
        "token_secret": token_secret,
        "project_name": os.environ.get("TABLEAU_PROJECT_NAME", "Bike-Sharing-Project"),
        "workbook_name": os.environ.get("TABLEAU_WORKBOOK_NAME", "Project_data_Analysis"),
    }


def publish_to_tableau_cloud(hyper_path=APP_HYPER_PATH):
    """Publish the app's latest extract by overwriting its existing Cloud workbook.

    Returns None when publishing is not configured. The workbook template must be the
    version whose dashboard layout should be retained on Tableau Cloud.
    """
    config = tableau_cloud_config()
    if config is None:
        return None

    try:
        import tableauserverclient as TSC
    except ImportError as exc:
        raise RuntimeError(
            "Tableau Cloud publishing requires tableauserverclient; install requirements.txt."
        ) from exc

    twbx_path = package_app_hyper(hyper_path)
    server = TSC.Server(config["server_url"], use_server_version=True)
    auth = TSC.PersonalAccessTokenAuth(
        config["token_name"],
        config["token_secret"],
        site_id=config["site_content_url"],
    )

    server.auth.sign_in(auth)
    try:
        projects = [
            project for project in TSC.Pager(server.projects)
            if project.name == config["project_name"]
        ]
        if len(projects) != 1:
            raise RuntimeError(
                f"Expected one Tableau project named {config['project_name']!r}; "
                f"found {len(projects)}."
            )
        project = projects[0]

        existing = [
            workbook for workbook in TSC.Pager(server.workbooks)
            if workbook.name == config["workbook_name"] and workbook.project_id == project.id
        ]
        if len(existing) != 1:
            raise RuntimeError(
                f"Expected one existing workbook named {config['workbook_name']!r} "
                f"in project {config['project_name']!r}; found {len(existing)}. "
                "Refusing to create or overwrite an ambiguous workbook."
            )

        workbook = TSC.WorkbookItem(project_id=project.id, name=config["workbook_name"])
        published = server.workbooks.publish(
            workbook,
            twbx_path,
            mode=TSC.Server.PublishMode.Overwrite,
        )
        return {"workbook_id": published.id, "workbook_name": published.name}
    finally:
        server.auth.sign_out()


def run_pipeline(excel_path=DEFAULT_EXCEL):
    """Runs steps (a) -> (b) -> (c) and returns a summary dict."""
    n = excel_to_hyper(excel_path)
    patch_twb()
    problems = validate_workbook()
    if problems:
        raise RuntimeError("; ".join(problems))
    package_twbx()
    return {
        "rows": n,
        "updated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "files": {"hyper": os.path.basename(HYPER_PATH), "twb": os.path.basename(TWB_PATH),
                  "twbx": os.path.basename(TWBX_PATH)},
    }


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Excel -> .hyper -> .twb/.twbx")
    ap.add_argument("--excel", default=DEFAULT_EXCEL)
    args = ap.parse_args()
    print(run_pipeline(args.excel))